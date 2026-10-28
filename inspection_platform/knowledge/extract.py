"""从字节流抽取纯文本，支持常见办公格式。"""

from __future__ import annotations

import io
import json
from pathlib import PurePath


def extract_plain_text(*, filename: str, content: bytes) -> tuple[str, str]:
    """
    返回 (format_key, plain_text)。
    format_key 为小写扩展名不含点，用于入库记录。
    """
    if not content:
        raise ValueError("empty file")
    ext = PurePath(filename).suffix.lower().lstrip(".")
    if ext in ("md", "markdown", "txt"):
        return ext, _decode_text(content)
    if ext == "json":
        return "json", _extract_json_text(content)
    if ext == "pdf":
        return "pdf", _extract_pdf(content)
    if ext in ("docx",):
        return "docx", _extract_docx(content)
    if ext in ("xlsx", "xlsm"):
        return ext, _extract_xlsx(content)
    if ext == "xls":
        raise ValueError("legacy .xls is not supported; please save as .xlsx")
    if ext in ("pptx",):
        return "pptx", _extract_pptx(content)
    if ext == "ppt":
        raise ValueError("legacy .ppt is not supported; please save as .pptx")
    raise ValueError(f"unsupported file extension: .{ext}")


def _decode_text(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "cp936"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_json_text(content: bytes) -> str:
    text = _decode_text(content)
    try:
        obj = json.loads(text)
    except json.JSONDecodeError as e:
        raise ValueError(f"invalid JSON: {e}") from e
    return json.dumps(obj, ensure_ascii=False, indent=2)


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    parts: list[str] = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        t = t.strip()
        if t:
            parts.append(t)
    return "\n\n".join(parts).strip()


def _extract_docx(content: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()


def _extract_xlsx(content: bytes) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in wb.worksheets:
            lines.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    lines.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(lines).strip()


def _extract_pptx(content: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    lines: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"## Slide {i}")
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            t = (shape.text or "").strip()
            if t:
                lines.append(t)
    return "\n".join(lines).strip()
