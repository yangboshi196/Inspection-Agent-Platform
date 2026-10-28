"""从 Excel 重新生成 ``builtin_seed.json``（需安装 openpyxl）。

用法（仓库根目录）::

    python -m inspection_platform.algorithm.seed_excel \"D:\\\\path\\\\算法场景评估测试-09_中广核_0926.xlsx\"
"""

from __future__ import annotations

import json
import sys
from pathlib import Path


def _norm(s: object) -> str:
    if s is None:
        return ""
    return str(s).strip()


def _health_from(s: str) -> str:
    if not s:
        return "unknown"
    if "较复杂" in s or "不好定义" in s:
        return "complex"
    if "需要确认" in s or "需要定义" in s:
        return "needs_definition"
    if s.startswith("可以检出"):
        return "ok"
    return "unknown"


def _build_schema(cat: str, seq: object, row_index: int) -> dict:
    hint: int | float | str
    if isinstance(seq, (int, float)):
        hint = int(seq) if isinstance(seq, float) and seq == int(seq) else seq
    else:
        hint = row_index
    return {
        "meta": {
            "source": "算法场景评估测试-09_中广核_0926.xlsx",
            "excel_row_hint": hint,
            "category": cat,
        },
        "input": {
            "type": "object",
            "properties": {
                "video_stream": {
                    "type": "string",
                    "description": "RTSP/HTTP 视频流地址或设备 id",
                },
                "roi": {"type": "array", "description": "可选 ROI 多边形顶点"},
                "time_window": {"type": "string", "description": "可选时间段配置"},
            },
        },
        "output": {
            "type": "object",
            "properties": {
                "detected": {"type": "boolean"},
                "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                "events": {"type": "array", "items": {"type": "object"}},
            },
        },
    }


def excel_to_records(xlsx_path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    hdr = list(rows[0])
    cols = {hdr[i]: i for i in range(len(hdr))}

    def get_cell(row: list[object], name: str) -> object:
        i = cols.get(name)
        if i is None:
            return None
        return row[i] if i < len(row) else None

    out: list[dict] = []
    counter = 0
    for ri, row in enumerate(rows[1:], start=2):
        row = list(row) + [""] * (len(hdr) - len(row))
        seq = get_cell(row, "序号")
        cat = _norm(get_cell(row, "算法场景类型"))
        scene = _norm(get_cell(row, "场景名称（字体黑色是√蓝色是x）"))
        app = _norm(get_cell(row, "算法应用说明"))
        std = _norm(get_cell(row, "检测标准要求"))
        prob = _norm(get_cell(row, "存在问题？"))
        model = _norm(get_cell(row, "模型是否可分析检出"))
        note = _norm(get_cell(row, "记录内容"))
        if not scene and not app:
            continue
        name = scene or (app[:48] + "..." if len(app) > 48 else app)
        if not name:
            continue
        counter += 1
        aid = f"algo_cgn_{counter:03d}"
        parts = [
            f"【场景类型】{cat}" if cat else "",
            f"【应用说明】{app}" if app else "",
            f"【检测标准要求】{std}" if std else "",
            f"【存在问题】{prob}" if prob else "",
            f"【模型可分析性】{model}" if model else "",
            f"【记录内容】{note}" if note else "",
        ]
        desc = "\n".join(p for p in parts if p)
        link = f"/api/algorithm/invoke/{aid}"
        out.append(
            {
                "algorithm_id": aid,
                "name": name,
                "desc": desc,
                "link": link,
                "schema": _build_schema(cat, seq, ri),
                "health": _health_from(model),
            }
        )
    return out


def main() -> None:
    if len(sys.argv) < 2:
        print(__doc__, file=sys.stderr)
        sys.exit(2)
    src = Path(sys.argv[1]).expanduser().resolve()
    if not src.is_file():
        print(f"文件不存在: {src}", file=sys.stderr)
        sys.exit(1)
    dst = Path(__file__).resolve().parent / "builtin_seed.json"
    recs = excel_to_records(src)
    dst.write_text(json.dumps(recs, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"写入 {dst}，共 {len(recs)} 条")


if __name__ == "__main__":
    main()
